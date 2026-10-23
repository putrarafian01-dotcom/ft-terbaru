import os, re, glob, logging
from pathlib import Path
from difflib import SequenceMatcher
from flask import Flask, request, jsonify
from flask_cors import CORS
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv('DATA_DIR', ROOT / 'data'))
EXCEL_PATH = os.getenv('EXCEL_PATH', '').strip()
ACCESS_CODE = os.getenv('ACCESS_CODE', '').strip()
TOP_K = max(1, min(int(os.getenv('TOP_K', '8')), 20))
MIN_SCORE = max(0.0, min(float(os.getenv('MIN_SCORE', '0.12')), 1.0))
logging.basicConfig(level=os.getenv('LOG_LEVEL', 'INFO'))
app = Flask(__name__)
CORS(app, resources={r'/api/*': {'origins': '*'}}, allow_headers=['Content-Type', 'X-Kode'], methods=['GET', 'POST', 'OPTIONS'])

QUESTION_KEYS = {'soal','pertanyaan','question','questions','nama soal','teks soal','query','q'}
ANSWER_KEYS = {'jawaban','jawab','answer','answers','kunci jawaban','jawaban benar','hasil','a'}
SOURCE_KEYS = {'sumber','source','file','nama file','sumber file'}

def norm(value):
    value = '' if value is None else str(value)
    value = value.lower().strip()
    value = re.sub(r'[^\w\s]', ' ', value, flags=re.UNICODE)
    return re.sub(r'\s+', ' ', value)

def header_key(value):
    return norm(value).replace('_', ' ')

def choose_column(headers, candidates):
    normalized = [header_key(h) for h in headers]
    for i, h in enumerate(normalized):
        if h in candidates:
            return i
    for i, h in enumerate(normalized):
        if any(c in h for c in candidates if len(c) > 3):
            return i
    return None

def load_rows():
    files = [Path(EXCEL_PATH)] if EXCEL_PATH and Path(EXCEL_PATH).is_file() else []
    if not files:
        files = [Path(p) for p in glob.glob(str(DATA_DIR / '**' / '*.xlsx'), recursive=True) if not Path(p).name.startswith('~$')]
        files += [Path(p) for p in glob.glob(str(DATA_DIR / '**' / '*.xlsm'), recursive=True) if not Path(p).name.startswith('~$')]
    rows, errors = [], []
    for path in files:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                iterator = ws.iter_rows(values_only=True)
                first = next(iterator, None)
                if not first: continue
                headers = [str(x or '').strip() for x in first]
                qi = choose_column(headers, QUESTION_KEYS)
                ai = choose_column(headers, ANSWER_KEYS)
                # Fallback: use first two columns if headers cannot be identified.
                if qi is None or ai is None:
                    qi, ai = 0, 1 if len(headers) > 1 else None
                if ai is None: continue
                si = choose_column(headers, SOURCE_KEYS)
                for row in iterator:
                    if not row or qi >= len(row) or ai >= len(row): continue
                    q, a = row[qi], row[ai]
                    if q is None or a is None or not str(q).strip() or not str(a).strip(): continue
                    source = row[si] if si is not None and si < len(row) and row[si] else path.name
                    rows.append({'q': str(q).strip(), 'a': str(a).strip(), 'sh': [str(source)], '_file': path.name})
            wb.close()
        except Exception as exc:
            errors.append(f'{path.name}: {exc}')
            app.logger.exception('Gagal membaca Excel %s', path)
    return rows, errors

def score(query, question):
    qn, sn = norm(query), norm(question)
    if not qn or not sn: return 0.0
    if qn == sn: return 1.0
    if qn in sn or sn in qn:
        ratio = min(len(qn), len(sn)) / max(len(qn), len(sn))
        return min(0.99, 0.82 + 0.17 * ratio)
    seq = SequenceMatcher(None, qn, sn).ratio()
    qt, st = set(qn.split()), set(sn.split())
    overlap = len(qt & st) / max(1, len(qt))
    jaccard = len(qt & st) / max(1, len(qt | st))
    return min(0.99, max(seq * 0.72 + overlap * 0.28, jaccard * 0.88))

@app.get('/')
def root():
    return jsonify({'name': 'Bank Soal Auto-Search API', 'status': 'ok', 'endpoints': ['/health', '/api/cari']})

@app.get('/health')
def health():
    rows, errors = load_rows()
    return jsonify({'ok': True, 'service': 'bank-soal-api', 'records': len(rows), 'excel_files_found': len(set(r['_file'] for r in rows)), 'excel_errors': errors}), 200

@app.route('/api/cari', methods=['POST', 'OPTIONS'])
def cari():
    if request.method == 'OPTIONS': return ('', 204)
    if ACCESS_CODE and request.headers.get('X-Kode', '') != ACCESS_CODE:
        return jsonify({'error': 'Kode akses tidak valid.'}), 401
    payload = request.get_json(silent=True) or {}
    query = str(payload.get('text', '')).strip()
    if not query:
        return jsonify({'c': [], 'error': 'Soal kosong. Masukkan teks soal.'}), 400
    if len(query) > 5000:
        return jsonify({'c': [], 'error': 'Teks soal terlalu panjang (maksimal 5000 karakter).'}), 413
    rows, errors = load_rows()
    if not rows:
        return jsonify({'c': [], 'error': 'Data Excel belum ditemukan atau tidak memiliki kolom soal dan jawaban.', 'excel_errors': errors}), 503
    results = []
    for row in rows:
        sim = score(query, row['q'])
        if sim >= MIN_SCORE:
            results.append({'q': row['q'], 'a': row['a'], 'sim': round(sim, 4), 'sh': row['sh']})
    results.sort(key=lambda item: item['sim'], reverse=True)
    return jsonify({'c': results[:TOP_K], 'total': len(results), 'source_count': len(rows)}), 200

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.getenv('PORT', '8000')), debug=False)
